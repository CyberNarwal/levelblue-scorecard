"""
Geometry QA for the walkthrough deck.

LibreOffice cannot run in this container, so the usual render-and-look pass is
unavailable. This checks the defects that pass would catch and that a reader
would see: a shape off the slide, a shape too close to the edge, two text boxes
overlapping, and text estimated to overflow its box.

The text-fit estimate is approximate - it assumes an average glyph width - so
it is tuned to flag gross overflow rather than to police the last line.
"""

import math
import sys
from pptx import Presentation
from pptx.util import Emu

EMU_PER_INCH = 914400
SLIDE_W = 13.333
SLIDE_H = 7.5
EDGE_MIN = 0.5

# Average glyph width as a fraction of point size, measured generously so the
# estimate errs towards reporting overflow rather than missing it.
WIDTH_RATIO = {"bold": 0.56, "normal": 0.51}
MONO_RATIO = 0.60


def inches(value):
    return None if value is None else value / EMU_PER_INCH


def run_info(shape):
    """Longest font size and whether any run is bold, across the shape."""
    size, bold, mono = 12.0, False, False
    if not shape.has_text_frame:
        return size, bold, mono
    for para in shape.text_frame.paragraphs:
        for run in para.runs:
            if run.font.size is not None:
                size = max(size, run.font.size.pt)
            if run.font.bold:
                bold = True
            if run.font.name and "Courier" in run.font.name:
                mono = True
    return size, bold, mono


def estimate_lines(text, width_in, size_pt, bold, mono):
    if not text.strip():
        return 0
    ratio = MONO_RATIO if mono else WIDTH_RATIO["bold" if bold else "normal"]
    chars_per_line = max(1, int((width_in * 72) / (size_pt * ratio)))
    lines = 0
    for para in text.split("\n"):
        lines += max(1, math.ceil(len(para) / chars_per_line))
    return lines


def overlaps(a, b):
    ax, ay, aw, ah = a
    bx, by, bw, bh = b
    return ax < bx + bw and bx < ax + aw and ay < by + bh and by < ay + ah


def main(path):
    prs = Presentation(path)
    problems = []

    for index, slide in enumerate(prs.slides, start=1):
        boxes = []
        for shape in slide.shapes:
            x, y = inches(shape.left), inches(shape.top)
            w, h = inches(shape.width), inches(shape.height)
            if None in (x, y, w, h):
                continue
            name = (shape.text_frame.text[:44].replace("\n", " ")
                    if shape.has_text_frame and shape.text_frame.text else shape.shape_type)

            if x < -0.01 or y < -0.01 or x + w > SLIDE_W + 0.01 or y + h > SLIDE_H + 0.01:
                problems.append(f"slide {index}: off the slide - {name!r} at "
                                f"({x:.2f}, {y:.2f}) {w:.2f}x{h:.2f}")
            elif x < EDGE_MIN - 0.01 or y < EDGE_MIN - 0.01 \
                    or x + w > SLIDE_W - EDGE_MIN + 0.01 or y + h > SLIDE_H - EDGE_MIN + 0.01:
                problems.append(f"slide {index}: inside the {EDGE_MIN}\" margin - {name!r} at "
                                f"({x:.2f}, {y:.2f}) {w:.2f}x{h:.2f}")

            if shape.has_text_frame and shape.text_frame.text.strip():
                size, bold, mono = run_info(shape)
                lines = estimate_lines(shape.text_frame.text, w, size, bold, mono)
                needed = lines * size * 1.22 / 72
                if needed > h + 0.06:
                    problems.append(
                        f"slide {index}: text may overflow - {name!r} needs about "
                        f"{needed:.2f}\" for {lines} line(s) at {size:.0f}pt, box is {h:.2f}\"")
                boxes.append(((x, y, w, h), name))

        for i in range(len(boxes)):
            for j in range(i + 1, len(boxes)):
                if overlaps(boxes[i][0], boxes[j][0]):
                    problems.append(f"slide {index}: text boxes overlap - "
                                    f"{boxes[i][1]!r} and {boxes[j][1]!r}")

    print(f"{len(prs.slides)} slides checked")
    if problems:
        print(f"\n{len(problems)} to look at:")
        for p in problems:
            print("  -", p)
        return 1
    print("no geometry problems found")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1]))
