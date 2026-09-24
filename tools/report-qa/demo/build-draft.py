"""
Build the deck used for the live demo.

Every fault in it is planted on purpose and verified by the demo script, so the
presenter knows exactly what the tool will say before they drop the file on it.
The client is fictional and the file is marked DEMO DRAFT in three places, so it
can never be mistaken for real work.

    python3 tools/report-qa/demo/build-draft.py
"""

from pathlib import Path

from pptx import Presentation
from pptx.oxml.ns import nsdecls
from pptx.util import Pt

from lxml.etree import fromstring as parse_xml

ROOT = Path(__file__).resolve().parents[3]
OUT = ROOT / "Northwind Trading - Security Assessment (DEMO DRAFT).pptx"

# Each entry: title, bullets, speaker notes. The faults are deliberate.
SLIDES = [
    (
        "Security Assessment - Northwind Trading",
        [
            "Prepared for Northwind Trading Ltd",
            "Engagement reference: TBC",          # placeholder-text blocker
            "DEMO DRAFT - not a real deliverable",
        ],
        "Don't mention the day rate here - push them for more scope in Q3.",
    ),
    (
        "Executive Summary",
        [
            "The organization was assessed against PCI DSS v3.2.1",   # US spelling + retired version
            "Overall posture is rated Medium with a CVSS score of 9.1",  # band contradicts score
            "MFA is not enforced for administrators",                 # acronym never expanded
            "Remediation is on-going and will be prioritized in Q4",   # house style + US spelling
        ],
        "Numbers here are a guesstimate - confirm before issue.",
    ),
    (
        "Findings",
        [
            "The whitelist was reviewed against Azure AD",            # two non-canonical names
            "Log4J remains unpatched on three hosts",                 # non-canonical product name
            "Findings map to PCI DSS requirement 14.2",               # requirement past the end
            "Control AC-0 was cited as the access control baseline",  # not a real 800-53 control
        ],
        "Chargeable extra if they ask for a retest.",
    ),
    (
        "Recommendations",
        [
            "It is important to note that in order to reduce risk, the organisation "
            "should very quickly implement multi factor authentication across all of "
            "the systems that are currently in scope for this particular engagement",
            "Review the e-mail gateway configuration",
            "Escalation paths are undocumented and should be defined",
        ],
        "",
    ),
]


def build():
    prs = Presentation()

    # The previous engagement's name, left in the slide master the way a real
    # template leak happens - invisible to anyone reading the slides. The shape
    # is built as XML because python-pptx cannot add one to a master, and a
    # scratch slide would leave an orphaned part behind in the package.
    leak = parse_xml(
        f'<p:sp {nsdecls("p", "a")}>'
        '<p:nvSpPr><p:cNvPr id="990" name="Footer"/>'
        '<p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr>'
        '<p:spPr><a:xfrm><a:off x="274320" y="6309360"/>'
        '<a:ext cx="4114800" cy="274320"/></a:xfrm>'
        '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/></p:spPr>'
        '<p:txBody><a:bodyPr wrap="none"/><a:lstStyle/><a:p><a:r>'
        '<a:rPr lang="en-GB" sz="900"/>'
        '<a:t>Contoso Financial Services - Confidential</a:t>'
        '</a:r></a:p></p:txBody></p:sp>'
    )
    prs.slide_master.shapes._spTree.append(leak)

    for title, bullets, notes in SLIDES:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        frame = slide.placeholders[1].text_frame
        frame.text = bullets[0]
        for bullet in bullets[1:]:
            frame.add_paragraph().text = bullet
        if notes:
            slide.notes_slide.notes_text_frame.text = notes

    prs.core_properties.title = "Security Assessment - DEMO DRAFT"
    prs.core_properties.author = "Contoso Financial Services"   # metadata leak
    prs.core_properties.comments = "DEMO FILE - fictional client, planted faults"
    prs.save(OUT)
    print(f"Built {OUT}")


if __name__ == "__main__":
    build()
